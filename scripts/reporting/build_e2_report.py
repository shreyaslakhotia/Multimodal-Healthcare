from pathlib import Path
import shutil
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import argparse
parser=argparse.ArgumentParser()
parser.add_argument('--input-dir',type=Path,required=True)
parser.add_argument('--output-dir',type=Path,required=True)
args=parser.parse_args()
BASE=args.input_dir
OUT=args.output_dir
OUT.mkdir(exist_ok=True)
summary=pd.read_csv(BASE/'summary.csv')
folds=pd.read_csv(BASE/'fold_metrics.csv')
curves=pd.read_csv(BASE/'learning_curves.csv')
comparisons=pd.read_csv(BASE/'comparisons.csv')
ensemble=summary.loc[summary.estimator=='mean_probability_ensemble'].set_index('modality')
MODS=['context','vitals','labs','cxr']
LABEL={'context':'Context','vitals':'Vitals','labs':'Labs','cxr':'CXR'}
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
sec.header_distance=sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]
    st.font.name='Calibri'
    st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(12)
sec.header.paragraphs[0].text='ICU MULTIMODAL BENCHMARK     |     E2 EVALUATION'
sec.header.paragraphs[0].style='Caption'
footer=sec.footer.paragraphs[0]
footer.text='21 October 2026     •     Research evaluation                                    '
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
md=[]
def title(t):
    doc.add_paragraph(t,'Title'); md.append('# '+t+'\n')
def h(t,level=1):
    doc.add_heading(t,level); md.append('#'*(level+1)+' '+t+'\n')
def p(t):
    doc.add_paragraph(t); md.append(t+'\n')
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Light Shading Accent 1'
    t.autofit=False
    for i,name in enumerate(headers):
        t.rows[0].cells[i].text=name
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
    for index,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        for i,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[i])
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(4)
                para.paragraph_format.space_before=Pt(4)
                for run in para.runs:
                    run.font.size=Pt(9)
                    run.font.color.rgb=RGBColor(0,0,0)
                    if index==0: run.bold=True
            if index==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7E6E6'); c._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows)+'\n')
def fig(name,caption,width=6.8):
    doc.add_picture(str(BASE/f'{name}.png'),width=Inches(width))
    doc.add_paragraph(caption,'Caption')
    md.append(f'![{caption}](E2Figures/{name}.png)\n')
def page(): doc.add_page_break()
def f(v): return f'{v:.3f}'

title('Evaluation of E2 Unimodal Encoders')
doc.add_paragraph('First 48 hour ICU mortality benchmark','Subtitle')
p('Evaluation date 21 October 2026. This document evaluates the four independently trained E2 encoders before fusion. All 60 planned fits completed. Vitals have the strongest unimodal AUROC point estimate, while the CXR encoder adds no statistically resolved AUROC improvement over its logistic baseline. The results support proceeding to fixed fusion experiments without an encoder architecture sweep.')
h('Background')
p('The project asks which clinical modalities contribute to in-hospital mortality prediction among adults alive and still in ICU at a 48-hour landmark. E1 established fixed logistic and LightGBM reference models; E2 replaces each modality’s summary classifier with a learned 128-dimensional representation. These representations will initialize the E3 and E4 fusion models.')
p('The final matched cohort contains 6,225 patients and 705 in-hospital deaths, a prevalence of 11.33%. Each patient contributes one first ICU stay. EHR observations are restricted to the first 48 hours. Chest X-rays lack timing metadata, so their results are exploratory associations rather than evidence of prospective availability at the landmark. Report text is excluded.')
table(['Modality','Ensemble AUROC with 95% CI','Average precision'],[[LABEL[m],f"{f(ensemble.loc[m,'auroc'])} ({f(ensemble.loc[m,'auroc_lower'])}–{f(ensemble.loc[m,'auroc_upper'])})",f(ensemble.loc[m,'auprc'])] for m in MODS],[1.1,3.7,2])
h('Scope of this evaluation',2)
p('The primary display averages the three prespecified seed probabilities for each patient; it does not select the best seed. Individual seed scores and fold dispersion are also reported. Comparisons use the same patients and fixed folds as E1. A unimodal ranking does not establish the incremental contribution of a modality in fusion.')
p('Training job 215538 completed on one NVIDIA L40S in 15 minutes 44 seconds, following a 19-second GPU smoke test. Evaluation job 217724 completed on CPU in 1 minute 22 seconds after 46 tests passed. There were no new prespecified leakage-gate failures. The existing missingness warning remains active.')

page(); h('Detailed Experiment Methodology')
h('Cohort and partitioning',2)
p('Starting from MIMIC-IV 2.1 and the supplied Kaggle CXR metadata, the pipeline selects adult first ICU stays, excludes inconsistent outcome metadata, and requires ICU residence and survival through hour 48. The selected frontal image must resolve uniquely and decode. Eligible stays require valid heart-rate and numeric laboratory observations. Final value QC removed no additional eligible stays.')
p('Five StratifiedGroupKFold outer folds are persisted with seed 42. There is one stay per patient and zero patient overlap within a fold’s training, validation and test partitions. Fifteen percent of the four-fold development portion is reserved for stratified inner validation, using seed 42 plus the fold number. No outer test patient contributes to preprocessing, checkpoint or threshold selection. Outer-fold death counts are 136, 158, 125, 141 and 145.')
h('Inputs and preprocessing',2)
table(['Stream','Input construction'],[
('Context','Age standardized on inner training patients; binary sex; training-only one-hot admission type and first care unit, with explicit unknown categories. Input width follows the fitted vocabulary rather than a hard-coded 20.'),
('Vitals','48 hourly bins × 16 channels: 8 normalized values and 8 observation masks. Values: heart rate, respiratory rate, SpO2, systolic/diastolic/mean blood pressure, temperature and total GCS.'),
('Labs','48 hourly bins × 22 channels: 11 normalized values and 11 masks. Glucose, pH, hemoglobin, white blood cells, platelets, creatinine, BUN, sodium, potassium, bicarbonate and lactate.'),
('CXR','49 × 1,024 frozen spatial DenseNet features. First listed AP image, otherwise first PA; no substitute if the selected file is missing. Grayscale, short side 224, center crop 224 and intensity scaling to [−1,024, 1,024].')],[.9,5.9])
p('For EHR, use observations in [ICU admission, admission plus 48 hours), hourly means, fixed unit conversion and broad artifact bounds in configs/features.yaml. GCS is assembled only when all three valid components occur within the same hour; ETT verbal responses are unassessable. Standardization statistics use observed inner-training values only. Forward fill stays within the patient window; leading gaps use the training mean.')
p('CXR features come from frozen torchxrayvision densenet121-res224-chex weights trained on CheXpert. The image backbone is not fine-tuned in E2. No PCA is used in E2; E1’s image comparator uses pooled features and training-fold PCA to 64 dimensions. Thus comparisons include representation and learner differences.')

page(); h('Detailed Experiment Methodology',1)
h('Fixed encoder architectures',2)
table(['Encoder','Architecture','Output before pooling'],[
('Context','MLP: input → 64 → 128; ReLU and dropout 0.1 between linear layers.','1 × 128'),
('Vitals','Linear 16 → 128; sinusoidal positions; 2 Transformer encoder layers; 4 heads; feed-forward width 256; dropout 0.1.','48 × 128'),
('Labs','Linear 22 → 128; otherwise the same temporal architecture as vitals.','48 × 128'),
('CXR','Linear 1,024 → 128; learned spatial positions; 1 Transformer encoder layer; 4 heads; feed-forward width 256; dropout 0.1.','49 × 128')],[.8,4.5,1.5])
p('Each model applies masked mean pooling followed by a single 128-to-1 linear mortality head. A temporal hour is pooled only if at least one channel was observed. All context/image tokens are valid. Temporal self-attention has no observation key-padding mask: it can attend to imputed tokens, although entirely unobserved hours are excluded from the final pool. Masks are also explicit input channels. This is the implemented E2 method, not an assumption of complete missingness invariance.')
table(['Training choice','Fixed value'],[
('Runs','4 modalities × 5 folds × seeds 0, 1 and 2 = 60'),
('Objective','BCE with logits; positive weight = training negatives / training positives'),
('Optimizer','AdamW; learning rate 0.0001; weight decay 0.01'),
('Batch and precision','Batch size 64; bf16 autocast; gradient norm clipped at 1.0'),
('Selection','At most 50 epochs; patience 10; checkpoint with highest inner-validation AUROC'),
('Inference','Best selected checkpoint; sigmoid probabilities; no test-based epoch or seed selection'),
('Resume','Per-epoch model, optimizer and RNG checkpoints; separate sampler/worker generators; saved per-fold encoders')],[1.45,5.35])
h('Evaluation and uncertainty',2)
p('For each modality, average corresponding seed probabilities before patient bootstrapping. Ensemble thresholds are selected on the similarly averaged validation predictions within each fold, then applied unchanged to that fold’s test patients. The main operating point maximizes validation F1, breaking ties at the highest threshold. Separate thresholds target validation recall of 0.80 and 0.90; actual test recall is reported without assuming those targets generalize.')
p('AUROC measures ranking; AUPRC is average precision, with a cohort prevalence reference of 0.113. The 95% intervals use 1,000 patient bootstrap draws, seed 42. Paired E2-minus-E1 differences share resampled patients. Intervals are conditional on the trained models and exclude uncertainty from retraining, alternate splits and external sites. Comparisons are descriptive and not multiplicity-adjusted. Three seed copies are never counted as three independent patients.')

page(); h('Results and Discussion')
h('Discrimination and seed stability',2)
fig('performance','Figure 1. Ensemble AUROC and average precision with 95% patient bootstrap intervals. Crosses show the three individual seed scores, not confidence limits. All results use the same 6,225 patients and five folds.')
table(['Modality','Seed AUROCs 0  1  2','Mean ± SD across seeds','Ensemble AP with 95% CI'],[[LABEL[m], ' / '.join(f(x) for x in summary.loc[(summary.modality==m)&(summary.estimator!='mean_probability_ensemble'),'auroc']),f"{summary.loc[(summary.modality==m)&(summary.estimator!='mean_probability_ensemble'),'auroc'].mean():.3f} ± {summary.loc[(summary.modality==m)&(summary.estimator!='mean_probability_ensemble'),'auroc'].std():.3f}",f"{f(ensemble.loc[m,'auprc'])} ({f(ensemble.loc[m,'auprc_lower'])}–{f(ensemble.loc[m,'auprc_upper'])})"] for m in MODS],[.8,2.15,1.65,2.2])
p('Vitals have the highest ensemble AUROC point estimate, followed by labs, context and CXR. Vitals and labs have similar average-precision point estimates with overlapping patient intervals. These are descriptive rankings: the analysis does not include a paired vitals-versus-labs significance test and does not establish which modality adds the most information after conditioning on the others.')
p('Averaging seeds improves the vitals and labs point estimates relative to each individual run. CXR shows the largest seed-to-seed AUROC spread. Seed standard deviations describe optimization sensitivity on these fixed folds; they are not patient confidence intervals or evidence of external generalization.')

page(); h('Results and Discussion')
fig('roc_pr','Figure 2. Locked pooled out-of-fold ROC and precision recall curves for the three-seed ensembles. The PR reference is mortality prevalence, not the random AUROC diagonal. CXR is untimed.')
h('Comparison with E1 counterparts',2)
chosen=comparisons.loc[comparisons.baseline.str.contains('counts1') | comparisons.modality.isin(['context','cxr'])]
rows=[]
for _,r in chosen.iterrows():
    learner='Logistic' if 'logistic' in r.baseline else 'LightGBM'
    counts=' + counts' if 'counts1' in r.baseline else ''
    rows.append([LABEL[r.modality],learner+counts,f"{r.auroc_delta:+.3f} ({r.auroc_lower:+.3f} to {r.auroc_upper:+.3f})",f"{r.auprc_delta:+.3f} ({r.auprc_lower:+.3f} to {r.auprc_upper:+.3f})"])
table(['Modality','E1 comparator','Delta AUROC with 95% CI','Delta AP with 95% CI'],rows,[.7,1.4,2.35,2.35])
p('E2 is not clearly worse than its relevant E1 counterpart. In particular, the vitals and labs AUROC intervals against count-aware tabular comparators include zero. E2 improves some no-count reference models, but E2 itself contains observation masks, so this cannot be attributed solely to temporal architecture. CXR improves over LightGBM, while its difference from logistic regression is unresolved. Context is also unresolved against logistic regression. There is no evidence here to justify an encoder architecture sweep.')
p('The table compares a three-seed E2 ensemble with the fixed seed-0 E1 model, not equal-sized ensembles of identical representations. Count-aware E1 models are emphasized because E2 explicitly receives masks. Full comparisons, including both count variants, are retained in E2Comparisons.csv. These comparisons do not show a multimodal gain; E2 models each modality independently.')

page(); h('Results and Discussion')
h('Training behavior and execution',2)
fig('learning_curves','Figure 3. Fifteen individual validation trajectories per modality, one per fold and seed. Curves are not seed-averaged, and end when early stopping or the epoch limit is reached. These are validation scores, not test learning curves.',width=6.55)
rows=[]
for m in MODS:
    frows=folds.loc[folds.modality==m]
    runs=curves.loc[curves.modality==m].groupby(['fold','seed']).epoch.max()
    rows.append([LABEL[m],f"{(frows.selected_epoch+1).median():.0f} [{(frows.selected_epoch+1).min():.0f}–{(frows.selected_epoch+1).max():.0f}]",f"{runs.median():.0f} [{runs.min():.0f}–{runs.max():.0f}]",f"{frows.auroc.mean():.3f} ± {frows.auroc.std():.3f}"])
table(['Modality','Selected epoch median range','Trained epochs median range','Fold and seed AUROC mean ± SD'],rows,[.8,1.8,1.8,2.4])
p('Context improves gradually, with some runs reaching the 50-epoch cap. Temporal models improve quickly before noisier validation plateaus; the CXR runs stop relatively early. The highest validation checkpoint, rather than the final epoch, supplies test predictions. Early stopping limits unnecessary computation without choosing among architectures. Fold-and-seed SD is descriptive because the 15 fits share patients and overlapping training sets.')
p('All fits completed within 15 minutes 44 seconds on a single L40S. CPU aggregation took 1 minute 22 seconds. The quick runtime reflects the small cohort, frozen image backbone, cached features and early stopping; it is not an estimate for larger cohorts or end-to-end image fine-tuning.')

page(); h('Results and Discussion')
h('Operating points and error patterns',2)
fig('confusion','Figure 4. Confusion counts for all 6,225 held-out patients per modality, using fold-specific validation maximum-F1 thresholds. Rows are actual labels and columns are predicted labels.')
table(['Modality','Precision','Recall','F1','Specificity'],[[LABEL[m]]+[f(ensemble.loc[m,k]) for k in ['precision','recall','f1','specificity']] for m in MODS],[1.1,1.4,1.4,1.4,1.5])
table(['Modality','Target 80 test precision recall','Target 90 test precision recall'],[[LABEL[m],f"{ensemble.loc[m,'validation_recall80_precision']:.3f} / {ensemble.loc[m,'validation_recall80_recall']:.3f}",f"{ensemble.loc[m,'validation_recall90_precision']:.3f} / {ensemble.loc[m,'validation_recall90_recall']:.3f}"] for m in MODS],[1,2.9,2.9])
p('Recall targets refer to validation, while the table gives the resulting test precision and recall. None of these operating points has been selected for clinical deployment. With 11.33% mortality prevalence, false positives remain substantial even when discrimination exceeds chance. Maximum F1 prioritizes a particular precision–recall trade-off and is not a clinical utility assessment. Calibration and decision-curve analysis were outside this experiment.')
h('Limits that govern interpretation',2)
p('The cohort is selected by a surviving 48-hour ICU landmark and available CXR, so results do not cover early deaths, earlier discharges or all ICU patients. It is a single-center retrospective sample. Unknown image timing may introduce future or cross-admission information. The final mask-only LightGBM AUROC remains 0.687, above the 0.65 warning threshold; measurement intensity may carry severity or care-process information. Neither moderate scores nor a passing shuffled-label canary proves absence of leakage. Clinical usefulness requires separate prospective, temporal and external validation.')

page(); h('Conclusion')
p('E2 is technically complete: 60 fits, saved fold-specific encoders and 12 complete seed-by-modality out-of-fold prediction files. Vitals and labs provide the strongest unimodal discrimination in this analysis. The corresponding ensemble AUROCs are 0.773 and 0.745; context is 0.684 and untimed CXR is 0.615. These estimates are compatible with the count-aware tabular counterparts and do not motivate an architecture search.')
p('The results establish reusable unimodal starting points, not a fusion benefit. They support the next fixed experiments while preserving the E1 reference models, locked folds, train-only preprocessing, active missingness warning and CXR timing caveat.')
h('Next steps')
p('Proceed with E3 early fusion: concatenate the four pooled 128-dimensional E2 representations, freeze the encoders and train a 512-to-128-to-1 MLP. Then use the same head with the E2 encoders fine-tuned jointly. Run all five folds and three seeds with the same fixed training settings, validation-only checkpoint/threshold selection and protected outer tests. The CheXpert DenseNet image backbone remains frozen; joint training concerns the E2 encoders above it.')
p('Add an unweighted average of E2 probabilities as a low-cost late-fusion reference. Logistic stacking is lower priority: its fitting and threshold selection need separate validation roles or a training-only cross-fitting design. Do not silently train and assess a stacker on the same validation predictions. The optional MedFuse-style LSTM remains lowest priority.')
p('After E3, compare paired patient AUROC and average-precision differences against E1 EHR, E1 all modalities and the late-fusion reference. Follow with E4 cross-attention and E5 masking/leave-one-out experiments to test incremental contribution and robustness. A unimodal rank or attention weight alone will not be presented as modality importance.')
h('Reproducibility and evidence',2)
table(['Evidence','Location or identifier'],[
('Training','Slurm 215538; source 0f6bd8d; E2 completion signature e1ed74ed723d556a3dc05899eba2b92423e2fa3162377bc3dd25e1bf9fcc0c23'),
('Evaluation','Slurm 217724; source 27f9b56; 46 tests passed; evaluate_e2.py and scripts/slurm/evaluate_e2.sbatch'),
('Private model artifacts','/projects/icumultimodal/results/deep/e2 and results/oof; patient predictions stay on the cluster'),
('Aggregate evidence','results/evaluation/e2/summary.csv, comparisons.csv, fold_metrics.csv, learning_curves.csv and metadata.json'),
('Figures','/projects/icumultimodal/reports/figures/e2; copied alongside this document'),
('Protocol and environment','docs/PROTOCOL.md; configs/features.yaml; requirements.txt; results/setup/pipfreeze.txt and condaexplicit.txt')],[1.15,5.65])
p('Every reported numerical result traces to the locked training/evaluation artifacts above. Exported CSVs and figures contain aggregate information only. No patient rows, images, report text, identifiers or predictions were used in this document authoring session.')

doc.save(OUT/'E2Evaluation.docx')
(OUT/'E2Evaluation.md').write_text('\n'.join(md),encoding='utf-8')
(OUT/'E2Figures').mkdir(exist_ok=True)
for name in ['performance','roc_pr','learning_curves','confusion']:
    shutil.copy2(BASE/f'{name}.png',OUT/'E2Figures'/f'{name}.png')
for source,target in [('summary.csv','E2Results.csv'),('comparisons.csv','E2Comparisons.csv'),('fold_metrics.csv','E2FoldMetrics.csv')]:
    shutil.copy2(BASE/source,OUT/target)
print('E2_REPORT_CREATED',OUT/'E2Evaluation.docx')
